import pandas as pd
import openpyxl


#Can be used to read out a table from an excel file by its name
def read_out_table(workbook_path, table_name, **dataframe_constructor_kwargs):
    wb = openpyxl.load_workbook(workbook_path, data_only=True)

    ws = next(sheet for sheet in wb.worksheets if table_name in sheet.tables.keys())
    table = ws.tables[table_name]
    rows = []
    for row in ws[table.ref]:
        rows.append([cell.value for cell in row])
    df = pd.DataFrame(rows[1:], columns=rows[0], **dataframe_constructor_kwargs)
    return df